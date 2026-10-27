"""services/export_excel.py - Xuất các biểu mẫu ra Excel (.xlsx) theo bố cục Hướng dẫn 07-HD/BTCTU.

Bố cục bám theo các tệp mẫu «4. Mau 2,4,5,6,7,8.xlsx» và «6. Mau 1, Mau 3 (so, nganh).xlsx»:
quốc hiệu, tiêu đề, số hiệu mẫu, tiêu đề cột nhiều tầng (ô gộp), dòng số hiệu cột (1)…(n), dòng trục/nhóm,
dòng tổng, khối ký xác nhận. Phông Times New Roman, kẻ khung, tự xuống dòng, khổ A4 ngang, lặp dòng tiêu đề
khi in. Mỗi hàm trả về (tên tệp, BytesIO).
"""
from __future__ import annotations

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import Axis, CommonCriterion, KpiCoefficient, Rating, ResponseLevel, TaskGroup

ORG = "ĐẢNG ỦY CÁC CƠ QUAN ĐẢNG\nTHÀNH PHỐ CẦN THƠ\n*"
NATION = "ĐẢNG CỘNG SẢN VIỆT NAM\nCần Thơ, ngày …… tháng …… năm ……"
FONT = "Times New Roman"
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREY = PatternFill("solid", fgColor="F2F2F2")
AXIS_FILL = PatternFill("solid", fgColor="FDE9E7")
RATINGS = dict(Rating.CHOICES)
GROUPS = dict(TaskGroup.CHOICES)
ROMAN = ["I", "II", "III", "IV"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def period_tag(period) -> str:
    return f"Q{period.quarter}-{period.year}"


def period_text(period) -> str:
    return f"QUÝ {ROMAN[period.quarter - 1]} NĂM {period.year}"


# --------------------------------------------------------------------------- #
# Tiện ích dựng trang tính
# --------------------------------------------------------------------------- #
class Sheet:
    def __init__(self, wb, title, widths, landscape=True):
        self.ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet()
        self.ws.title = title
        self.n = len(widths)
        for i, w in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        ps = self.ws.page_setup
        ps.paperSize = self.ws.PAPERSIZE_A4
        ps.orientation = "landscape" if landscape else "portrait"
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        ps.fitToWidth, ps.fitToHeight = 1, 0
        self.ws.page_margins.left = self.ws.page_margins.right = 0.4
        self.ws.page_margins.top = self.ws.page_margins.bottom = 0.5
        self.row = 1

    def cell(self, r, c, value=None, bold=False, italic=False, size=12, align="left", valign="center",
             border=False, fill=None, fmt=None, wrap=True):
        cl = self.ws.cell(row=r, column=c, value=value)
        cl.font = Font(name=FONT, size=size, bold=bold, italic=italic)
        cl.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
        if border:
            cl.border = BOX
        if fill:
            cl.fill = fill
        if fmt:
            cl.number_format = fmt
        return cl

    def merge(self, r1, c1, r2, c2, value=None, border=False, **kw):
        self.cell(r1, c1, value, border=border, **kw)
        if (r1, c1) != (r2, c2):
            self.ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
            if border:
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        self.ws.cell(row=r, column=c).border = BOX

    def heading(self, form_no, title_lines, left_span=None, org=ORG):
        """Quốc hiệu (trái: tên cơ quan; phải: Đảng Cộng sản Việt Nam), tiêu đề, số hiệu mẫu."""
        half = left_span or max(2, self.n // 2 - 1)
        self.merge(1, 1, 1, half, org, bold=True, align="center", size=12)
        self.merge(1, half + 1, 1, self.n, NATION, bold=True, align="center", size=12)
        self.ws.row_dimensions[1].height = 48
        self.merge(3, 1, 3, self.n, "\n".join(title_lines), bold=True, align="center", size=13)
        self.ws.row_dimensions[3].height = 18 * len(title_lines) + 6
        self.cell(4, self.n, form_no, bold=True, italic=True, align="right")
        self.row = 5

    def header_cell(self, r1, c1, r2, c2, text):
        self.merge(r1, c1, r2, c2, text, border=True, bold=True, align="center", size=11, fill=GREY)

    def colnos(self, labels, row=None):
        r = row or self.row
        for i, lab in enumerate(labels, start=1):
            self.cell(r, i, lab, border=True, italic=True, align="center", size=10)
        self.row = r + 1
        return r

    def line(self, values, bold=False, italic=False, fill=None, fmts=None, aligns=None, height=None):
        r = self.row
        for i, v in enumerate(values, start=1):
            fmt = (fmts or {}).get(i)
            align = (aligns or {}).get(i, "right" if isinstance(v, (int, float)) and not isinstance(v, bool) else "left")
            self.cell(r, i, v, border=True, bold=bold, italic=italic, fill=fill, fmt=fmt, align=align,
                      valign="top", size=11)
        if height:
            self.ws.row_dimensions[r].height = height
        self.row += 1
        return r

    def band(self, first, text, fill=AXIS_FILL, bold=True, italic=False):
        """Dòng trục / nhóm: ô đầu ghi nhãn ngắn, phần còn lại gộp ghi tên."""
        r = self.row
        self.cell(r, 1, first, border=True, bold=bold, fill=fill, align="center", size=11)
        self.merge(r, 2, r, self.n, text, border=True, bold=bold, italic=italic, fill=fill, size=11)
        self.row += 1
        return r

    def para(self, text, bold=False, italic=False, height=None, c1=1, c2=None):
        r = self.row
        self.merge(r, c1, r, c2 or self.n, text, bold=bold, italic=italic, valign="top")
        if height:
            self.ws.row_dimensions[r].height = height
        self.row += 1
        return r

    def signatures(self, left=None, right=None, gap=1):
        self.row += gap
        r = self.row
        half = self.n // 2
        if left:
            self.merge(r, 1, r, half, left, bold=True, align="center", valign="top")
        if right:
            self.merge(r, half + 1, r, self.n, right, bold=True, align="center", valign="top")
        self.ws.row_dimensions[r].height = 64
        self.row += 6

    def repeat_header(self, first, last):
        self.ws.print_title_rows = f"{first}:{last}"
        self.ws.freeze_panes = self.ws.cell(row=last + 1, column=1)


def _done(wb, filename):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return filename, buf


def _axes():
    return Axis.query.order_by(Axis.number).all()


def _num(v):
    return None if v is None else float(v)


def _dmy(d):
    if not d:
        return ""
    if isinstance(d, str):
        y, m, dd = d.split("-")
        return f"{dd}/{m}/{y}"
    return d.strftime("%d/%m/%Y")


# --------------------------------------------------------------------------- #
# Mẫu 1, Mẫu 2 - KPI tập thể
# --------------------------------------------------------------------------- #
def mau1(ce, indicators, leaders):
    wb = Workbook()
    widths = [7, 46, 16, 11, 11, 24, 18, 18, 20, 18]
    sh = Sheet(wb, "Mau 1", widths)
    sh.heading("Mẫu 1", ["DANH MỤC CHỈ TIÊU, NHIỆM VỤ, SẢN PHẨM CẦN THỰC HIỆN",
                         f"{ce.department.name.upper()} - {period_text(ce.period)}"])
    labels = ["STT", "Nội dung, nhiệm vụ, chỉ tiêu", "Sản phẩm / Đơn vị tính", "Chỉ tiêu năm", "Chỉ tiêu quý",
              "Độ khó, mới, phức tạp; phạm vi tác động", "Cơ quan chủ trì", "Cơ quan phối hợp",
              "Lãnh đạo trực tiếp phụ trách", "Thời gian hoàn thành"]
    top = sh.row
    for i, lab in enumerate(labels, start=1):
        sh.header_cell(top, i, top, i, lab)
    sh.ws.row_dimensions[top].height = 48
    sh.row += 1
    sh.colnos([f"({i})" for i in range(1, 11)])
    sh.repeat_header(top, top + 1)
    lead = {u["id"]: u["name"] for u in leaders}
    for axis in _axes():
        rows = [i for i in indicators if i.axis_id == axis.id]
        if not rows:
            continue
        sh.band(f"Trục {axis.number}", axis.name)
        for key, title in TaskGroup.CHOICES:
            g = [i for i in rows if i.task_group == key]
            if not g:
                continue
            sh.band("", title, fill=None, italic=True)
            for n, ind in enumerate(g, start=1):
                sh.line([n, ind.content, " / ".join(x for x in (ind.product, ind.unit) if x),
                         _num(ind.annual_target), _num(ind.quarter_target), ind.difficulty or "",
                         ind.lead_agency or "", ind.support_agency or "", lead.get(ind.leader_id, ""),
                         " · ".join(x for x in (ind.deadline_text, _dmy(ind.deadline_date)) if x)],
                        aligns={1: "center"})
    sh.signatures(None, "XÁC NHẬN CỦA TẬP THỂ LÃNH ĐẠO CƠ QUAN, ĐƠN VỊ\n"
                        "(Xác lập thời điểm, ký, ghi rõ họ tên và đóng dấu)")
    return _done(wb, f"Mau1_{ce.department.code}_{period_tag(ce.period)}.xlsx")


def mau2(ce, indicators):
    wb = Workbook()
    widths = [6, 26, 44, 10, 12, 12, 12, 12, 12, 16, 14]
    sh = Sheet(wb, "Mau 2", widths)
    sh.heading("Mẫu 2", [f"KẾT QUẢ THỰC HIỆN NHIỆM VỤ CỦA {ce.department.name.upper()} {period_text(ce.period)}"],
               left_span=3)
    sh.para("I. KẾT QUẢ THỰC HIỆN NHIỆM VỤ THEO CHỈ TIÊU NGHỊ QUYẾT NĂM", bold=True)
    labels = ["STT", "Trục/Nhóm\nnhiệm vụ", "Chỉ tiêu/\nNội dung nhiệm vụ/Sản phẩm đầu ra", "Đơn vị tính",
              f"Chỉ tiêu nhiệm vụ năm {ce.period.year}", f"Chỉ tiêu nhiệm vụ quý {ROMAN[ce.period.quarter - 1]}",
              f"Kết quả thực hiện quý {ROMAN[ce.period.quarter - 1]}", "Tỷ lệ (%) so với chỉ tiêu quý",
              "Tỷ lệ (%) so với chỉ tiêu năm", "Thời gian hoàn thành theo quy định (nếu có)",
              "Thời gian hoàn thành thực tế"]
    top = sh.row
    for i, lab in enumerate(labels, start=1):
        sh.header_cell(top, i, top, i, lab)
    sh.ws.row_dimensions[top].height = 60
    sh.row += 1
    sh.colnos(["(1)", "(2)", "(3)", "(4)", "(5)", "(6)", "(7)", "(8)=(7)/(6)*100", "(9)=(7)/(5)*100",
               "(10)", "(11)"])
    sh.repeat_header(top, top + 1)
    no = 0
    for axis in _axes():
        for key, title in TaskGroup.CHOICES:
            g = [i for i in indicators if i.axis_id == axis.id and i.task_group == key]
            if not g:
                continue
            first = sh.row
            for ind in g:
                no += 1
                r = ind.collective_kpi
                sh.line([no, None, ind.content, ind.unit or "", _num(ind.annual_target), _num(ind.quarter_target),
                         _num(r.result) if r else None, _num(r.pct_quarter) if r else None,
                         _num(r.pct_year) if r else None,
                         " · ".join(x for x in (ind.deadline_text, _dmy(ind.deadline_date)) if x),
                         _dmy(r.actual_date) if r else ""],
                        fmts={8: "0.00", 9: "0.00"}, aligns={1: "center"})
            sh.merge(first, 2, sh.row - 1, 2, f"Trục {axis.number}. {axis.name}\n{title}", border=True,
                     bold=True, valign="top", size=11)
    sh.row += 1
    sh.para("II. KẾT QUẢ THỰC HIỆN MỘT SỐ NHIỆM VỤ TRỌNG TÂM KHÁC", bold=True)
    sh.para(ce.other_key_tasks or "(Do đơn vị tổng hợp theo tình hình thực hiện nhiệm vụ thực tế)",
            italic=not ce.other_key_tasks, height=max(30, 15 * (1 + (ce.other_key_tasks or "").count("\n"))))
    sh.para("III. MÔ HÌNH HAY, CÁCH LÀM SÁNG TẠO (nếu có)", bold=True)
    sh.para(ce.innovative_models or "", height=max(30, 15 * (1 + (ce.innovative_models or "").count("\n"))))
    sh.signatures(None, "XÁC NHẬN CỦA TẬP THỂ LÃNH ĐẠO CƠ QUAN, ĐƠN VỊ\n"
                        "(Xác lập thời điểm, ký, ghi rõ họ tên và đóng dấu)")
    return _done(wb, f"Mau2_{ce.department.code}_{period_tag(ce.period)}.xlsx")


# --------------------------------------------------------------------------- #
# Mẫu 3, Mẫu 4, Mẫu 5 - KPI cá nhân
# --------------------------------------------------------------------------- #
def _owner_line(ev):
    u = ev.user
    return f"Họ và tên: {u.full_name}    Chức vụ: {ev.position_title or u.role_label}    " \
           f"Đơn vị: {u.department.name if u.department else ''}"


def mau3(ev, plan_rows):
    wb = Workbook()
    widths = [12, 46, 12, 26, 18, 10, 18, 12, 11]
    sh = Sheet(wb, "Mau 3", widths)
    sh.heading("Mẫu 3", ["DANH MỤC SẢN PHẨM/CÔNG VIỆC CỦA CÁ NHÂN CẦN THỰC HIỆN",
                         f"Quý {ROMAN[ev.period.quarter - 1]}/{ev.period.year}"])
    sh.para(_owner_line(ev), italic=True)
    labels = ["Trục/Nhóm nhiệm vụ", "Sản phẩm/công việc", "Cấp trình", "Độ khó, mới, phức tạp; phạm vi tác động",
              "Sản phẩm", "Số lượng", "Tiến độ\n(ghi cụ thể ngày hoàn thành nếu có)", "Điểm chấm công việc\ncụ thể",
              "Hệ số quy đổi"]
    top = sh.row
    for i, lab in enumerate(labels, start=1):
        sh.header_cell(top, i, top, i, lab)
    sh.ws.row_dimensions[top].height = 48
    sh.row += 1
    sh.colnos([f"({i})" for i in range(1, 10)])
    sh.repeat_header(top, top + 1)
    for axis in _axes():
        rows = [r for r in plan_rows if r["axis_id"] == axis.id]
        if not rows:
            continue
        sh.band(f"Trục {axis.number}", axis.name)
        for gi, (key, title) in enumerate(TaskGroup.CHOICES, start=1):
            g = [r for r in rows if r["task_group"] == key]
            if not g:
                continue
            sh.band(f"Nhóm {gi}", title, fill=None, italic=True)
            for r in g:
                sh.line(["", r["task"], "", r.get("difficulty") or "", r.get("product") or "", _num(r.get("planned_qty")),
                         " · ".join(x for x in (r.get("deadline_text"), _dmy(r.get("deadline_date"))) if x),
                         _num(r.get("task_score")), _num(r.get("conversion_coef"))],
                        fmts={9: "0.00"})
    sh.signatures(None, "CÁ NHÂN LẬP DANH MỤC\n(Ký, ghi rõ họ tên)")
    return _done(wb, f"Mau3_{ev.user.username}_{period_tag(ev.period)}.xlsx")


def mau4(ev):
    wb = Workbook()
    widths = [5, 36, 14, 9, 14, 9, 10, 8, 8, 8, 8, 12, 8, 8, 9, 8, 8, 8, 8]
    sh = Sheet(wb, "Mau 4", widths)
    sh.heading("Mẫu 4", [f"NHIỆM VỤ ĐỀ RA KẾ HOẠCH VÀ THỰC TẾ TRIỂN KHAI {period_text(ev.period)}"], left_span=5)
    sh.para(_owner_line(ev), italic=True)
    t = sh.row
    sh.header_cell(t, 1, t + 2, 1, "TT")
    sh.header_cell(t, 2, t, 6, "NHIỆM VỤ ĐỀ RA KẾ HOẠCH")
    sh.header_cell(t, 7, t, 16, "THỰC TẾ TRIỂN KHAI")
    sh.header_cell(t, 17, t + 1, 19, "Hệ số KPI vượt")
    for c, lab in zip(range(2, 7), ["Nhiệm vụ", "Sản phẩm", "Số lượng", "Tiến độ", "Hệ số quy đổi"]):
        sh.header_cell(t + 1, c, t + 2, c, lab)
    sh.header_cell(t + 1, 7, t + 1, 9, "KPI\n(Số lượng)")
    sh.header_cell(t + 1, 10, t + 1, 11, "KPI\n(Chất lượng)")
    sh.header_cell(t + 1, 12, t + 1, 14, "KPI\n(Tiến độ)")
    sh.header_cell(t + 1, 15, t + 1, 16, "Lãnh đạo, chỉ đạo, điều hành")
    for c, lab in zip(range(7, 20), ["Số lượng thực tế hoàn thành", "Hệ số KPI", "Quy đổi", "Hệ số KPI", "Quy đổi",
                                     "Thời điểm thực tế hoàn thành", "Hệ số KPI", "Quy đổi", "Hệ số KPI", "Quy đổi",
                                     "KPI (Số lượng)", "KPI (Chất lượng)", "KPI (Tiến độ)"]):
        sh.header_cell(t + 2, c, t + 2, c, lab)
    sh.ws.row_dimensions[t + 1].height = 32
    sh.ws.row_dimensions[t + 2].height = 46
    sh.row = t + 3
    sh.colnos([f"({i})" for i in range(1, 20)])
    sh.repeat_header(t, t + 3)
    lead = ev.use_leadership
    f2 = {i: "0.00" for i in (6, 8, 9, 10, 11, 13, 14, 15, 16, 17, 18, 19)}

    def sums(items):
        s = lambda k: round(sum(getattr(i, k) or 0 for i in items), 2)
        return [s("conversion_coef"), None, None, s("qty_conv"), None, s("quality_conv"), None, None,
                s("progress_conv"), None, s("leadership_conv") if lead else None,
                s("over_qty"), s("over_quality"), s("over_progress")]

    no = 0
    results = {ar.axis_id: ar for ar in ev.axis_results}
    for axis in _axes():
        rows = [i for i in ev.items if i.axis_id == axis.id]
        if not rows:
            continue
        sh.band(f"Trục {axis.number}", axis.name)
        for key, title in TaskGroup.CHOICES:
            g = [i for i in rows if i.task_group == key]
            if not g:
                continue
            sh.band("", title, fill=None, italic=True)
            for it in g:
                no += 1
                sh.line([no, it.task, it.product or "", _num(it.planned_qty),
                         " · ".join(x for x in (it.deadline_text, _dmy(it.deadline_date)) if x),
                         _num(it.conversion_coef), _num(it.actual_qty), _num(it.kpi_qty), _num(it.qty_conv),
                         _num(it.quality_value) if it.quality_code else None, _num(it.quality_conv),
                         _dmy(it.actual_date), _num(it.progress_value) if it.progress_code else None,
                         _num(it.progress_conv),
                         (_num(it.leadership_value) if it.leadership_code else None) if lead else "—",
                         _num(it.leadership_conv) if lead else "—",
                         _num(it.over_qty) or None, _num(it.over_quality) or None, _num(it.over_progress) or None],
                        fmts=f2, aligns={1: "center", 15: "center" if not lead else "right", 16: "center" if not lead else "right"})
        ar = results.get(axis.id)
        label = f"Cộng trục {axis.number}" + (f" · KPI {ar.kpi_total:.2f}%" if ar else "")
        r = sh.line(["", label, "", "", ""] + sums(rows), bold=True, fill=GREY, fmts=f2)
        sh.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    r = sh.line(["", "TỔNG", "", "", ""] + sums(ev.items), bold=True, fill=GREY, fmts=f2)
    sh.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    sh.signatures("XÁC NHẬN CỦA TẬP THỂ LÃNH ĐẠO CƠ QUAN, ĐƠN VỊ\n(Xác lập thời điểm, ký, ghi rõ họ tên và đóng dấu)",
                  "CÁ NHÂN TỰ ĐÁNH GIÁ\n(Ký, ghi rõ họ tên)")
    return _done(wb, f"Mau4_{ev.user.username}_{period_tag(ev.period)}.xlsx")


def mau5(ev):
    wb = Workbook()
    widths = [7, 30, 18, 30, 14, 14, 11, 16, 20]
    sh = Sheet(wb, "Mau 5", widths, landscape=False)
    sh.heading("Mẫu 5", ["BẢN TỰ ĐÁNH GIÁ, XẾP LOẠI CỦA CÁ NHÂN",
                         f"Quý {ROMAN[ev.period.quarter - 1]}, Năm {ev.period.year}"], left_span=4)
    u = ev.user
    for text in (f"Họ và tên: {u.full_name}    Ngày sinh: {_dmy(u.birth_date)}",
                 f"Chức vụ Đảng: {u.party_position or ''}",
                 f"Chức vụ chính quyền: {u.gov_position or ev.position_title or u.role_label}",
                 f"Chức vụ đoàn thể: {u.org_position or ''}",
                 f"Đơn vị công tác: {u.department.name if u.department else ''}"):
        sh.para(text, c1=2)
    sh.para("I. Tự đánh giá kết quả thực hiện nhiệm vụ", bold=True, c1=2)
    sh.para("Trên cơ sở nhiệm vụ được giao, cá nhân tự đánh giá về kết quả thực hiện nhiệm vụ theo quý như sau:", c1=2)
    # --- A ---
    r = sh.row
    sh.cell(r, 1, "A", bold=True, border=True, align="center")
    sh.merge(r, 2, r, 9, "NHÓM TIÊU CHÍ CHUNG (30 ĐIỂM) - Các tiêu chí thực hiện theo Quy định số 89-QĐ/TU của "
                         "Ban Thường vụ Thành ủy", border=True, bold=True)
    sh.row += 1
    t = sh.row
    sh.header_cell(t, 1, t, 1, "TT")
    sh.header_cell(t, 2, t, 4, "Tiêu chí / Nội dung")
    for c, lab in zip(range(5, 10), ["Đảm bảo\n(Đánh dấu x)", "Không\nđảm bảo\n(Đánh dấu x)", "Điểm tối đa",
                                     "Điểm đạt (Chấm tối đa nếu đảm bảo; Chấm 0 điểm nếu không đảm bảo)", "Ghi chú"]):
        sh.header_cell(t, c, t, c, lab)
    sh.ws.row_dimensions[t].height = 62
    sh.row += 1
    r = sh.row
    for c, lab in [(1, "(1)"), (2, "(2)"), (5, "(3)"), (6, "(4)"), (7, "(5)"), (8, "(6)"), (9, "(7)")]:
        sh.cell(r, c, lab, border=True, italic=True, align="center", size=10)
    sh.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    sh.row += 1
    scores = {cs.criterion_id: cs for cs in ev.common_scores}
    crit = CommonCriterion.query.order_by(CommonCriterion.sort).all()
    last = None
    for c in crit:
        if c.group_no != last:
            last = c.group_no
            members = [x for x in crit if x.group_no == c.group_no]
            got = sum((scores[x.id].score or 0) for x in members if x.id in scores)
            r = sh.line([str(c.group_no), c.group_title, None, None, None, None,
                         sum(x.max_score for x in members), got, None], bold=True, aligns={1: "center"})
            sh.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        cs = scores.get(c.id)
        sat = cs.satisfied if cs else None
        r = sh.line([c.code, c.title, None, None, "x" if sat is True else "", "x" if sat is False else "",
                     c.max_score, (cs.score if cs and sat is not None else None), (cs.note if cs else "") or ""],
                    aligns={1: "center", 5: "center", 6: "center"})
        sh.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        sh.ws.row_dimensions[r].height = max(30, 15 * (len(c.title) // 60 + 1))
    r = sh.line(["Tổng (A) =", None, None, None, None, None, 30, ev.common_score, None], bold=True, fill=GREY,
                fmts={8: "0.00"})
    sh.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    # --- B ---
    r = sh.row
    sh.cell(r, 1, "B", bold=True, border=True, align="center")
    sh.merge(r, 2, r, 9, "KẾT QUẢ THỰC HIỆN NHIỆM VỤ ĐƯỢC GIAO (70 ĐIỂM)", border=True, bold=True)
    sh.row += 1
    t = sh.row
    sh.header_cell(t, 1, t, 1, "TT")
    sh.header_cell(t, 2, t, 3, "Tiêu chí / Nội dung")
    for c, lab in zip(range(4, 10), ["Mục tiêu, nhiệm vụ đề ra", "Kết quả sản phẩm thực tế", "Điểm KPI (%)",
                                     "Điểm tối đa", "Điểm đạt\n(= Điểm KPI x Điểm tối đa)/100", "Ghi chú"]):
        sh.header_cell(t, c, t, c, lab)
    sh.ws.row_dimensions[t].height = 46
    sh.row += 1
    axes = {a.id: a for a in _axes()}
    for n, ar in enumerate(sorted(ev.axis_results, key=lambda a: axes[a.axis_id].number), start=1):
        a = axes[ar.axis_id]
        r = sh.line([n, f"Trục {a.number}. {a.name}", None, ar.goals or "", ar.results or "", ar.kpi_total,
                     ar.max_score, ar.score, ar.note or ""], fmts={6: "0.00", 7: "0.00", 8: "0.00"},
                    aligns={1: "center"})
        sh.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        lines = max((ar.goals or "").count("\n"), (ar.results or "").count("\n")) + 2
        sh.ws.row_dimensions[r].height = min(400, max(45, 15 * lines))
    total_max = round(sum(ar.max_score for ar in ev.axis_results), 2)
    r = sh.line(["TỔNG (B) = ", None, None, None, None, None, total_max, ev.result_score, None], bold=True,
                fill=GREY, fmts={7: "0.00", 8: "0.00"})
    sh.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r = sh.line(["TỔNG (A+ B) = ", None, None, None, None, None, 30 + total_max, ev.total_score, None], bold=True,
                fill=GREY, fmts={7: "0.00", 8: "0.00"})
    sh.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    sh.row += 1
    sh.para(f"II. Tự đề xuất xếp loại mức chất lượng: {RATINGS.get(ev.self_rating, '')}", bold=True, c1=2)
    sh.para("(Theo 04 mức: 1- Hoàn thành xuất sắc nhiệm vụ, 2- Hoàn thành tốt nhiệm vụ, 3- Hoàn thành nhiệm vụ và "
            "4- Không hoàn thành nhiệm vụ)", italic=True, c1=2)
    sh.signatures(None, "CÁ NHÂN TỰ ĐÁNH GIÁ\n(Ký, ghi rõ họ tên)", gap=0)
    sh.para("III. Nhận xét, đánh giá của cấp có thẩm quyền", bold=True, c1=2)
    levels = dict(ResponseLevel.CHOICES)
    sh.para(f"- Chấm điểm: {ev.reviewer_score if ev.reviewer_score is not None else ''}", c1=2)
    sh.para(f"- Đề xuất xếp loại: {RATINGS.get(ev.final_rating or ev.proposed_rating, '')}", c1=2)
    sh.para("- Mức độ đáp ứng đối với các mục tiêu, nhiệm vụ then chốt: "
            f"{levels.get(ev.response_level, '')}", c1=2)
    if ev.reviewer_comment:
        sh.para(f"- Nhận xét: {ev.reviewer_comment}", c1=2)
    sh.signatures(None, "XÁC NHẬN CỦA TẬP THỂ LÃNH ĐẠO CƠ QUAN, ĐƠN VỊ\n"
                        "(Xác lập thời điểm, ký, ghi rõ họ tên và đóng dấu)")
    return _done(wb, f"Mau5_{ev.user.username}_{period_tag(ev.period)}.xlsx")


# --------------------------------------------------------------------------- #
# Mẫu 6, 7, 8 - báo cáo tổng hợp
# --------------------------------------------------------------------------- #
def mau6(rows, period):
    wb = Workbook()
    sh = Sheet(wb, "Mau 6", [6, 26, 34, 20, 20, 46, 30])
    sh.heading("Mẫu 6", [f"TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ VÀ ĐỀ XUẤT XẾP LOẠI QUÝ {ROMAN[period.quarter - 1]}, NĂM {period.year}",
                         "ĐỐI VỚI CÁN BỘ THUỘC DIỆN BAN THƯỜNG VỤ THÀNH ỦY ĐÁNH GIÁ, XẾP LOẠI CHẤT LƯỢNG"], left_span=3)
    labels = ["STT", "Họ và tên", "Chức vụ, đơn vị công tác", "Cá nhân tự đề xuất mức xếp loại",
              "Cấp có thẩm quyền đề xuất mức xếp loại",
              "Tóm tắt căn cứ, cơ sở lý do trong trường hợp đề xuất mức xếp loại hoàn thành xuất sắc hoặc mức xếp "
              "loại không hoàn thành nhiệm vụ hoặc các nội dung khác (nếu có)",
              "Đề xuất nội dung liên quan về công tác cán bộ (nếu có)"]
    top = sh.row
    for i, lab in enumerate(labels, start=1):
        sh.header_cell(top, i, top, i, lab)
    sh.ws.row_dimensions[top].height = 76
    sh.row += 1
    sh.colnos([f"({i})" for i in range(1, 8)])
    sh.repeat_header(top, top + 1)
    for n, r in enumerate(rows, start=1):
        sh.line([n, r["name"], r["position"], RATINGS.get(r["self"], ""), RATINGS.get(r["proposed"], ""),
                 r["reason"], r["personnel"]], aligns={1: "center"})
    sh.signatures(None, "XÁC NHẬN CỦA TẬP THỂ LÃNH ĐẠO CƠ QUAN, ĐƠN VỊ\n"
                        "(Xác lập thời điểm, ký, ghi rõ họ tên và đóng dấu)")
    return _done(wb, f"Mau6_{period_tag(period)}.xlsx")


NOTE_89 = ("Theo Điểm 1.5, Khoản 1, Điều 11, Quy định số 89-QĐ/TU: \"Số lượng tập thể, cá nhân được xếp loại "
           "“Hoàn thành xuất sắc nhiệm vụ” không vượt quá 20% tổng số tập thể, cá nhân được xếp loại “Hoàn thành tốt "
           "nhiệm vụ” theo đối tượng có vị trí, vai trò, chức năng, nhiệm vụ tương đồng, do cấp có thẩm quyền xem xét, "
           "quyết định\". Do đó, cấp có thẩm quyền khi đánh giá, xếp loại sẽ xem xét, quyết định các nhóm đối tượng có "
           "vị trí, vai trò, chức năng, nhiệm vụ tương đồng nhưng phải đảm bảo từng nhóm đối tượng có tối thiểu từ 5 "
           "cán bộ trở lên.")


def mau78(data, period, form):
    groups = data["groups"]
    n = len(groups)
    wb = Workbook()
    sh = Sheet(wb, "Mau 7" if form == "m7" else "Mau 8", [6, 34, 14] + [16] * n + [24])
    scope = "CẤP XÃ VÀ TƯƠNG ĐƯƠNG" if form == "m7" else "CẤP THÀNH PHỐ"
    sh.heading("Mẫu 7" if form == "m7" else "Mẫu 8",
               [f"BÁO CÁO KẾT QUẢ ĐÁNH GIÁ CÁN BỘ LÃNH ĐẠO, QUẢN LÝ {scope}", period_text(period)], left_span=3)
    labels = ["STT", "Kết quả xếp loại", "Tổng số cán bộ lãnh đạo, quản lý"] + [g.name for g in groups] + \
             ["Ghi chú (ghi cụ thể đối tượng cán bộ lãnh đạo, quản lý khác)"]
    top = sh.row
    for i, lab in enumerate(labels, start=1):
        sh.header_cell(top, i, top, i, lab)
    sh.ws.row_dimensions[top].height = 110
    sh.row += 1
    sh.colnos(["(1)", "(2)", "(3)=" + "+".join(f"({i})" for i in range(4, n + 4))] +
              [f"({i})" for i in range(4, n + 4)] + [f"({n + 4})"])
    sh.repeat_header(top, top + 1)
    for row in data["rows"]:
        if row.get("ratio"):
            vals = [row["ratio"]["total"]] + [row["ratio"]["per"][g.id] for g in groups]
        elif row.get("cells"):
            vals = [row["cells"]["total"]] + [row["cells"]["per"][g.id] for g in groups]
        else:
            vals = [None] * (n + 1)
        sh.line([row["no"], row["label"]] + vals + [None], bold=row.get("bold", False), italic=row.get("italic", False),
                fmts={i: "0.00" for i in range(3, n + 4)} if row.get("ratio") else None, aligns={1: "center"})
    sh.row += 1
    sh.para(("Lưu ý: " if form == "m7" else "Ghi chú: ") + NOTE_89, italic=True, height=75)
    sh.signatures("NGƯỜI LẬP BIỂU\n(Ký, ghi rõ họ tên, số điện thoại để liên hệ)",
                  ("XÁC NHẬN CỦA BAN THƯỜNG VỤ ĐẢNG UỶ XÃ, PHƯỜNG, ĐẢNG UỶ TRỰC THUỘC THÀNH ỦY" if form == "m7"
                   else "XÁC NHẬN CỦA TẬP THỂ LÃNH ĐẠO CƠ QUAN, ĐƠN VỊ") +
                  "\n(Xác lập thời điểm, ký, ghi rõ họ tên và đóng dấu)")
    return _done(wb, f"{'Mau7' if form == 'm7' else 'Mau8'}_{period_tag(period)}.xlsx")
