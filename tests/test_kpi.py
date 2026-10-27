"""Kiểm thử KPI cá nhân 3 bước (Mẫu 3 -> 4 -> 5), KPI tập thể, luồng trình - duyệt, công thức Phụ lục 2
và nâng cấp CSDL.

Người dùng mẫu (vai trò mới):
  leader    Phó Bí thư Thường trực (lãnh đạo cơ quan)   manager1  Phó TB Thường trực, Phòng 1
  leader_p1 Phó Trưởng Ban, Phòng 1                     manager2  Phó TB Thường trực, Phòng 2
  staff1    Công chức, Phòng 1                          staff2    Công chức, Phòng 2
"""
import copy

import pytest
from sqlalchemy import text

from app import create_app
from migrate import upgrade_schema
from models import (
    Axis, CollectiveEvaluation, CommonCriterion, Department, EvalStatus, Indicator, Period,
    PersonalEvaluation, PersonalKPIItem, PlanStatus, Role, User, WorkflowLog, db,
)
from seed import seed_all
from services.auth import create_user

PW = "CanTho2026kpi"

# Ví dụ minh họa trong Phụ lục 2 (Bảng 1, trục xây dựng Đảng):
# (nhiệm vụ, kế hoạch, thực tế, điểm công việc, chất lượng, tiến độ, lãnh đạo)
EXAMPLE = [
    ("Báo cáo tháng", 3, 2, 100, 1, 3, 2),
    ("Báo cáo quý", 1, 1, 100, 2, 1, 2),
    ("Báo cáo tình hình cơ sở đảng", 1, 1, 120, 3, 1, 2),
    ("Tờ trình chi bộ bốn tốt", 1, 1, 150, 1, 1, 1),
    ("Đề án nâng cao năng lực lãnh đạo", 1, 1, 200, 1, 1, 1),
    ("Kết nạp đảng viên mới", 20, 25, 120, 1, 1, 1),
]


@pytest.fixture
def app():
    app = create_app("testing")
    with app.app_context():
        db.create_all()
        seed_all()
        root = Department(code="TU", name="Thành ủy")
        p1 = Department(code="P1", name="Phòng 1")
        p2 = Department(code="P2", name="Phòng 2")
        db.session.add_all([root, p1, p2])
        db.session.flush()
        for username, role, dept in [
            ("admin", Role.ADMIN, None), ("leader", Role.PBT, root.id),
            ("leader_p1", Role.PTB, p1.id), ("manager1", Role.PTBTT, p1.id), ("manager2", Role.PTBTT, p2.id),
            ("staff1", Role.CC, p1.id), ("staff2", Role.CC, p2.id),
        ]:
            create_user(username, username.title(), PW, role=role, department_id=dept, approved=True)
        db.session.add(Period.build(2026, 3))
        db.session.add(Period.build(2026, 4))
        db.session.commit()
    yield app
    with app.app_context():
        db.session.remove()
        db.drop_all()


def client(app, username):
    c = app.test_client()
    assert c.post("/login", data={"username": username, "password": PW}).status_code == 302
    return c


def ids(app):
    with app.app_context():
        return {
            "axis": {a.number: a.id for a in Axis.query.all()},
            "criteria": [c.id for c in CommonCriterion.query.order_by(CommonCriterion.sort)],
            "q3": Period.query.filter_by(year=2026, quarter=3).first().id,
            "q4": Period.query.filter_by(year=2026, quarter=4).first().id,
            "p1": Department.query.filter_by(code="P1").first().id,
            "p2": Department.query.filter_by(code="P2").first().id,
            "tu": Department.query.filter_by(code="TU").first().id,
        }


def uid(app, username):
    with app.app_context():
        return User.query.filter_by(username=username).first().id


def plan_payload(app, rows=EXAMPLE, group="thuong_xuyen", axis_no=4):
    """Bước 1 · Mẫu 3: chỉ phần danh mục."""
    axis = ids(app)["axis"][axis_no]
    return {"items": [{"axis_id": axis, "task_group": group, "task": t, "planned_qty": pl, "task_score": sc}
                      for (t, pl, ac, sc, q, p, l) in rows]}


def full_payload(app, rows=EXAMPLE, max_score=70, rating="HTT", satisfied=True, group="thuong_xuyen",
                 item_ids=None, axis_no=4):
    """Bước 2, 3 · Mẫu 4, 5. item_ids: id các dòng đã có (theo thứ tự rows)."""
    i = ids(app)
    axis = i["axis"][axis_no]
    item_ids = list(item_ids or [])
    return {
        "self_rating": rating, "self_comment": "Tự đánh giá",
        "axes": [{"axis_id": axis, "is_main": True, "max_score": max_score}],
        "common": [{"criterion_id": c, "satisfied": satisfied} for c in i["criteria"]],
        "items": [{
            "id": item_ids[n] if n < len(item_ids) else None,
            "axis_id": axis, "task_group": group, "task": t, "planned_qty": pl, "actual_qty": ac,
            "task_score": sc, "quality_code": q, "progress_code": p, "leadership_code": l,
        } for n, (t, pl, ac, sc, q, p, l) in enumerate(rows)],
    }


def start(app, c, period="q4"):
    r = c.post(f"/kpi/start/{ids(app)[period]}")
    assert r.status_code == 302
    return int(r.headers["Location"].rsplit("/", 1)[1])


def ev_of(app, eid):
    with app.app_context():
        ev = db.session.get(PersonalEvaluation, eid)
        return {"status": ev.status, "plan": ev.plan_status, "items": [i.id for i in ev.items]}


def status(app, eid):
    return ev_of(app, eid)["status"]


def plan_status(app, eid):
    return ev_of(app, eid)["plan"]


def approve_plan(app, eid, who):
    """Đưa Mẫu 3 đã trình qua đủ chuỗi duyệt của người lập `who`."""
    ban_leader = {"staff1": "manager1", "leader_p1": "manager1", "staff2": "manager2"}.get(who)
    if ban_leader:
        assert client(app, ban_leader).post(f"/kpi/{eid}/plan/approve").status_code == 302
        assert plan_status(app, eid) == PlanStatus.BAN_APPROVED
    assert client(app, "leader").post(f"/kpi/{eid}/plan/approve").status_code == 302
    assert plan_status(app, eid) == PlanStatus.APPROVED


def ready(app, who="staff1", rows=EXAMPLE, period="q4", group="thuong_xuyen"):
    """Lập, trình, duyệt xong Mẫu 3 -> sẵn sàng tự chấm Mẫu 4. Trả về (client, eid, id các dòng)."""
    c = client(app, who)
    eid = start(app, c, period)
    assert c.post(f"/kpi/{eid}/save", json=plan_payload(app, rows, group)).status_code == 200
    if who == "leader":
        assert c.post(f"/kpi/{eid}/plan/lock").status_code == 302
        assert plan_status(app, eid) == PlanStatus.APPROVED
    else:
        assert c.post(f"/kpi/{eid}/plan/submit").status_code == 302
        assert plan_status(app, eid) == PlanStatus.SUBMITTED
        approve_plan(app, eid, who)
    return c, eid, ev_of(app, eid)["items"]


def to_submitted(app, who="staff1", rows=EXAMPLE):
    c, eid, item_ids = ready(app, who, rows)
    assert c.post(f"/kpi/{eid}/save", json=full_payload(app, rows, item_ids=item_ids)).status_code == 200
    assert c.post(f"/kpi/{eid}/submit").status_code == 302
    assert status(app, eid) == EvalStatus.SELF_SCORED
    return c, eid


APPROVE = {"final_rating": "HTT", "reviewer_comment": "Tốt"}


# --------------------------- công thức --------------------------- #
def test_formula_matches_phu_luc_2_example(app):
    c, eid, item_ids = ready(app, "leader_p1")                     # Phó Trưởng Ban: có chỉ số D
    r = c.post(f"/kpi/{eid}/save", json=full_payload(app, item_ids=item_ids))
    assert r.status_code == 200, r.get_json()
    d = r.get_json()["data"]
    ax = d["axes"][0]
    assert (ax["kpi_a"], ax["kpi_b"], ax["kpi_c"]) == (95.82, 95.70, 97.47)      # đúng số trong Phụ lục 2
    assert ax["kpi_d"] == 95.95 and ax["kpi_total"] == 96.24
    assert ax["score"] == 67.37 and ax["over_qty"] == 0.25                       # 25/20 -> vượt 0,25, không cộng vào KPI
    assert d["totals"] == {"common": 30.0, "result": 67.37, "total": 97.37}
    first = d["items"][0]
    assert first["kpi_qty"] == 0.67 and first["qty_conv"] == 0.67
    assert d["items"][2]["quality_conv"] == 0.96                                  # 0,8 x 1,2
    assert d["items"][5]["kpi_qty"] == 1.0                                        # vượt mức bị chặn ở 1


def test_cong_chuc_has_no_leadership_kpi(app):
    c, eid, item_ids = ready(app, "staff1")                         # công chức: KPI = (A+B+C)/3
    p = full_payload(app, item_ids=item_ids)
    p["use_leadership"] = True                                      # gửi lên cũng bị bỏ qua
    d = c.post(f"/kpi/{eid}/save", json=p).get_json()["data"]
    ax = d["axes"][0]
    assert d["use_leadership"] is False
    assert ax["kpi_d"] == 0 and ax["kpi_total"] == round((95.82 + 95.70 + 97.47) / 3, 2)
    data = page_data(c, eid)
    assert data["leadership"] is False
    _, e2, ids2 = ready(app, "leader_p1")
    assert page_data(client(app, "leader_p1"), e2)["leadership"] is True


def page_data(c, eid):
    import json, re
    html = c.get(f"/kpi/{eid}").get_data(as_text=True)
    return json.loads(re.search(r'<script id="kpi-data" type="application/json">(.*?)</script>', html, re.S).group(1))


def test_unanswered_common_criteria_score_zero(app):
    c, eid, item_ids = ready(app)
    d = c.post(f"/kpi/{eid}/save", json=full_payload(app, satisfied=None, item_ids=item_ids)).get_json()["data"]
    assert d["totals"]["common"] == 0
    d = c.post(f"/kpi/{eid}/save", json=full_payload(app, satisfied=False, item_ids=item_ids)).get_json()["data"]
    assert d["totals"]["common"] == 0


# --------------------------- Bước 1 · Mẫu 3 --------------------------- #
def test_plan_stage_saves_only_plan_fields(app):
    c = client(app, "staff1")
    eid = start(app, c)
    d = c.post(f"/kpi/{eid}/save", json=full_payload(app)).get_json()["data"]     # gửi cả kết quả cũng bị bỏ qua
    assert d["stage"] == "plan" and len(d["items"]) == 6
    assert all(i["actual_qty"] is None and i["quality_code"] is None for i in d["items"])
    assert d["totals"]["common"] == 0 and d["self_rating"] == ""
    assert c.post(f"/kpi/{eid}/submit").status_code == 302 and status(app, eid) == "draft"   # chưa duyệt Mẫu 3


def test_plan_chain_lock_return_reopen(app):
    c = client(app, "staff1")
    eid = start(app, c)
    c.post(f"/kpi/{eid}/save", json=plan_payload(app))
    assert client(app, "leader").post(f"/kpi/{eid}/plan/approve").status_code == 403   # chưa trình: không xem được
    assert plan_status(app, eid) == PlanStatus.DRAFT
    c.post(f"/kpi/{eid}/plan/submit")
    assert plan_status(app, eid) == PlanStatus.SUBMITTED
    assert c.post(f"/kpi/{eid}/save", json=plan_payload(app)).status_code == 409        # đã trình -> khóa
    client(app, "leader").post(f"/kpi/{eid}/plan/approve")                              # phải qua lãnh đạo Ban trước
    assert plan_status(app, eid) == PlanStatus.SUBMITTED
    assert client(app, "manager2").post(f"/kpi/{eid}/plan/approve").status_code == 403  # Ban khác
    m1 = client(app, "manager1")
    m1.post(f"/kpi/{eid}/plan/return", data={"note": ""})                               # thiếu lý do
    assert plan_status(app, eid) == PlanStatus.SUBMITTED
    m1.post(f"/kpi/{eid}/plan/return", data={"note": "Bổ sung sản phẩm"})
    assert plan_status(app, eid) == PlanStatus.DRAFT
    assert "Bổ sung sản phẩm" in c.get(f"/kpi/{eid}").get_data(as_text=True)
    c.post(f"/kpi/{eid}/plan/submit")
    approve_plan(app, eid, "staff1")
    with app.app_context():
        actions = [l.action for l in WorkflowLog.query.filter_by(entity_type="personal", entity_id=eid).order_by(WorkflowLog.id)]
        assert actions == ["plan_submit", "plan_return", "plan_submit", "plan_ban_approve", "plan_approve"]
    # chỉ Admin hoặc Phó Bí thư mở lại
    assert client(app, "manager1").post(f"/kpi/{eid}/plan/reopen").status_code == 403
    assert c.post(f"/kpi/{eid}/plan/reopen").status_code == 403
    client(app, "admin").post(f"/kpi/{eid}/plan/reopen")
    assert plan_status(app, eid) == PlanStatus.DRAFT
    assert c.post(f"/kpi/{eid}/save", json=plan_payload(app, EXAMPLE[:2])).status_code == 200   # sửa lại được
    c.post(f"/kpi/{eid}/plan/submit")
    approve_plan(app, eid, "staff1")
    client(app, "leader").post(f"/kpi/{eid}/plan/reopen")
    assert plan_status(app, eid) == PlanStatus.DRAFT


def test_ptbtt_plan_goes_straight_to_pbt(app):
    c = client(app, "manager1")
    eid = start(app, c)
    c.post(f"/kpi/{eid}/save", json=plan_payload(app))
    c.post(f"/kpi/{eid}/plan/submit")
    assert client(app, "manager2").post(f"/kpi/{eid}/plan/approve").status_code == 403
    client(app, "leader").post(f"/kpi/{eid}/plan/approve")
    assert plan_status(app, eid) == PlanStatus.APPROVED


def test_pbt_saves_without_submitting(app):
    c = client(app, "leader")
    eid = start(app, c)
    c.post(f"/kpi/{eid}/save", json=plan_payload(app))
    html = c.get(f"/kpi/{eid}").get_data(as_text=True)
    assert "Trình Mẫu 3" not in html and "Chốt Mẫu 3" in html
    c.post(f"/kpi/{eid}/plan/submit")                                               # không có bước trình
    assert plan_status(app, eid) == PlanStatus.DRAFT
    c.post(f"/kpi/{eid}/plan/lock")
    assert plan_status(app, eid) == PlanStatus.APPROVED
    item_ids = ev_of(app, eid)["items"]
    assert c.post(f"/kpi/{eid}/save", json=full_payload(app, item_ids=item_ids)).status_code == 200
    html = c.get(f"/kpi/{eid}").get_data(as_text=True)
    assert "Trình Mẫu 5" not in html and "Không cần trình" in html
    c.post(f"/kpi/{eid}/submit")
    assert status(app, eid) == "draft"
    assert "Không cần trình" in c.get("/").get_data(as_text=True)                   # dashboard
    assert page_data(c, eid)["leadership"] is True


def test_plan_snapshot_unchanged_by_mau4(app):
    c, eid, item_ids = ready(app)
    p = full_payload(app, item_ids=item_ids)
    p["items"][0]["task"] = "Đổi tên ở Mẫu 4"                       # nội dung Mẫu 3 không đổi được ở Mẫu 4
    p["items"][0]["planned_qty"] = 99
    p["items"][1]["task_score"] = 150                              # hệ số quy đổi điều chỉnh được
    p["items"].append({"axis_id": ids(app)["axis"][4], "task_group": "trong_tam", "task": "Việc mới phát sinh",
                       "planned_qty": 1, "actual_qty": 1, "task_score": 100, "quality_code": 1, "progress_code": 1})
    d = c.post(f"/kpi/{eid}/save", json=p).get_json()["data"]
    by_id = {i["id"]: i for i in d["items"]}
    assert by_id[item_ids[0]]["task"] == "Báo cáo tháng" and by_id[item_ids[0]]["planned_qty"] == 3
    assert by_id[item_ids[1]]["conversion_coef"] == 1.5
    new = [i for i in d["items"] if i["id"] not in item_ids]
    assert len(d["items"]) == 7 and new[0]["from_plan"] is False and new[0]["task"] == "Việc mới phát sinh"
    assert len(d["plan"]) == 6 and d["plan"][1]["conversion_coef"] == 1.0            # bản chụp Mẫu 3 giữ nguyên
    assert [r["task"] for r in d["plan"]] == [r[0] for r in EXAMPLE]
    p2 = full_payload(app, rows=EXAMPLE[1:], item_ids=item_ids[1:])               # xóa nhiệm vụ của Mẫu 3
    r = c.post(f"/kpi/{eid}/save", json=p2)
    assert r.status_code == 400 and "không được xóa" in " ".join(r.get_json()["errors"])


def test_saving_with_few_axes(app):
    c, eid, item_ids = ready(app, rows=[EXAMPLE[0]])               # chỉ 1 trục
    p = full_payload(app, rows=[EXAMPLE[0]], item_ids=item_ids)
    assert c.post(f"/kpi/{eid}/save", json=p).status_code == 200
    assert c.post(f"/kpi/{eid}/submit").status_code == 302 and status(app, eid) == EvalStatus.SELF_SCORED
    c.post(f"/kpi/{eid}/withdraw")
    p["axes"] = [{"axis_id": ids(app)["axis"][4], "max_score": 40}, {"axis_id": ids(app)["axis"][1], "max_score": 30}]
    c.post(f"/kpi/{eid}/save", json=p)                             # trục 1 có điểm nhưng không có nhiệm vụ
    html = c.post(f"/kpi/{eid}/submit", follow_redirects=True).get_data(as_text=True)
    assert "chưa có nhiệm vụ nào" in html and status(app, eid) == "draft"


# --------------------------- kiểm tra dữ liệu khi lưu --------------------------- #
def test_save_validation(app):
    c, eid, item_ids = ready(app)
    base = full_payload(app, item_ids=item_ids)
    axis = ids(app)["axis"][4]

    def bad(mutate):
        p = copy.deepcopy(base)
        mutate(p)
        r = c.post(f"/kpi/{eid}/save", json=p)
        assert r.status_code == 400, r.get_json()
        return r.get_json()["errors"]

    def new(**kw):
        row = {"axis_id": axis, "task_group": "thuong_xuyen", "task": "Mới", "planned_qty": 1, "task_score": 100}
        row.update(kw)
        return lambda p: p["items"].append(row)

    assert any("lớn hơn 70" in e for e in bad(lambda p: p["axes"][0].update(max_score=71)))
    two = ids(app)["axis"]
    assert any("vượt 70" in e for e in bad(lambda p: p.update(axes=[
        {"axis_id": two[4], "max_score": 40}, {"axis_id": two[1], "max_score": 40}])))
    bad(lambda p: p["items"][0].update(task_score=201))
    bad(lambda p: p["items"][0].update(quality_code=9))
    bad(lambda p: p["items"][0].update(actual_qty=-1))
    bad(lambda p: p["items"][0].update(actual_date="31/12/2026"))
    bad(new(task=""))
    bad(new(axis_id=99999))
    bad(new(planned_qty="abc"))
    bad(lambda p: p.update(self_rating="XUATSAC"))
    assert c.post(f"/kpi/{eid}/save", data="khong phai json").status_code == 400
    with app.app_context():                       # lưu lỗi không để lại dữ liệu
        assert PersonalKPIItem.query.count() == 6


def test_cannot_touch_other_peoples_items(app):
    c1, c2 = client(app, "staff1"), client(app, "staff2")
    e1, e2 = start(app, c1), start(app, c2)
    c1.post(f"/kpi/{e1}/save", json=plan_payload(app))
    with app.app_context():
        foreign = db.session.get(PersonalEvaluation, e1).items[0].id
    p = plan_payload(app)
    p["items"][0]["id"] = foreign
    assert c2.post(f"/kpi/{e2}/save", json=p).status_code == 400
    assert c2.post(f"/kpi/{e1}/save", json=plan_payload(app)).status_code == 403   # hồ sơ người khác


def test_save_updates_and_deletes_rows_in_plan(app):
    c = client(app, "staff1")
    eid = start(app, c)
    d = c.post(f"/kpi/{eid}/save", json=plan_payload(app)).get_json()["data"]
    keep = d["items"][:2]
    p = {"items": [dict(k, planned_qty=5) for k in keep]}
    d2 = c.post(f"/kpi/{eid}/save", json=p).get_json()["data"]
    assert [i["id"] for i in d2["items"]] == [k["id"] for k in keep]
    assert d2["items"][0]["planned_qty"] == 5


# --------------------------- Bước 3 · Mẫu 5: trình -> nhận xét -> duyệt --------------------------- #
def test_submit_requires_complete_profile(app):
    c, eid, item_ids = ready(app)
    p = full_payload(app, max_score=60, item_ids=item_ids)          # tổng điểm trục 60 != 70
    p["items"][0]["quality_code"] = None
    p["common"][0]["satisfied"] = None
    p["self_rating"] = None
    c.post(f"/kpi/{eid}/save", json=p)
    html = c.post(f"/kpi/{eid}/submit", follow_redirects=True).get_data(as_text=True)
    for fragment in ("chưa chọn mức chất lượng", "tiêu chí chung chưa chọn", "phải bằng 70", "tự đề xuất xếp loại"):
        assert fragment in html, fragment
    assert "chưa chọn mức lãnh đạo" not in html                                     # công chức: không có D
    assert status(app, eid) == "draft"


def test_full_flow_draft_submit_review_approve(app):
    c, eid = to_submitted(app)
    assert c.post(f"/kpi/{eid}/save", json=full_payload(app)).status_code == 409    # đã trình thì không sửa
    mgr = client(app, "manager1")
    r = mgr.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT", "reviewer_comment": "Ổn"}, follow_redirects=True)
    assert status(app, eid) == EvalStatus.REVIEWED, r.get_data(as_text=True)[:200]
    lead = client(app, "leader")
    lead.post(f"/kpi/{eid}/approve", data=APPROVE)
    assert status(app, eid) == EvalStatus.APPROVED
    with app.app_context():
        ev = db.session.get(PersonalEvaluation, eid)
        assert ev.final_rating == "HTT" and ev.proposed_rating == "HTT"
        assert ev.reviewed_by.username == "manager1" and ev.approved_by.username == "leader"
        actions = [l.action for l in WorkflowLog.query.filter_by(entity_type="personal", entity_id=eid).order_by(WorkflowLog.id)]
        assert actions[-3:] == ["submit", "review", "approve"]
    assert c.post(f"/kpi/{eid}/withdraw").status_code == 302 and status(app, eid) == EvalStatus.APPROVED
    assert c.post(f"/kpi/{eid}/save", json=full_payload(app)).status_code == 409


def test_leader_can_approve_directly(app):
    c, eid = to_submitted(app)
    client(app, "leader").post(f"/kpi/{eid}/approve", data=APPROVE)
    with app.app_context():
        ev = db.session.get(PersonalEvaluation, eid)
        assert ev.status == EvalStatus.APPROVED and ev.reviewed_by.username == "leader"


def test_review_and_approve_input_rules(app):
    c, eid = to_submitted(app)
    mgr = client(app, "manager1")
    mgr.post(f"/kpi/{eid}/review", data={})                                          # thiếu xếp loại
    mgr.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTXS"})                 # HTXS cần căn cứ
    mgr.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT", "force_majeure": "on"})   # thiếu lý do
    mgr.post(f"/kpi/{eid}/review", data={"proposed_rating": "XX"})
    assert status(app, eid) == EvalStatus.SELF_SCORED
    mgr.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTXS", "rating_reason": "Vượt mức"})
    assert status(app, eid) == EvalStatus.REVIEWED


def test_key_task_requires_response_level(app):
    rows = [("Việc then chốt", 1, 1, 100, 1, 1, 1)]
    c, eid, item_ids = ready(app, rows=rows, group="trong_tam")
    c.post(f"/kpi/{eid}/save", json=full_payload(app, rows, group="trong_tam", item_ids=item_ids))
    c.post(f"/kpi/{eid}/submit")
    mgr = client(app, "manager1")
    mgr.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT"})
    assert status(app, eid) == EvalStatus.SELF_SCORED
    mgr.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT", "response_level": "1"})
    assert status(app, eid) == EvalStatus.REVIEWED


def test_withdraw_and_return(app):
    c, eid = to_submitted(app)
    item_ids = ev_of(app, eid)["items"]
    assert c.post(f"/kpi/{eid}/withdraw").status_code == 302 and status(app, eid) == "draft"
    c.post(f"/kpi/{eid}/submit")
    mgr = client(app, "manager1")
    mgr.post(f"/kpi/{eid}/return", data={"note": ""})                                # thiếu lý do
    assert status(app, eid) == EvalStatus.SELF_SCORED
    mgr.post(f"/kpi/{eid}/return", data={"note": "Bổ sung minh chứng"})
    assert status(app, eid) == "draft"
    assert "Bổ sung minh chứng" in c.get(f"/kpi/{eid}").get_data(as_text=True)
    assert c.post(f"/kpi/{eid}/save", json=full_payload(app, item_ids=item_ids)).status_code == 200
    c.post(f"/kpi/{eid}/submit")
    assert status(app, eid) == EvalStatus.SELF_SCORED
    mgr.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT"})
    c.post(f"/kpi/{eid}/withdraw")
    assert status(app, eid) == EvalStatus.REVIEWED
    assert client(app, "admin").post(f"/kpi/{eid}/plan/reopen").status_code == 302      # Mẫu 5 đã trình
    assert plan_status(app, eid) == PlanStatus.APPROVED


def test_reopen_approved(app):
    c, eid = to_submitted(app)
    lead = client(app, "leader")
    lead.post(f"/kpi/{eid}/approve", data=APPROVE)
    assert client(app, "manager1").post(f"/kpi/{eid}/reopen").status_code == 403
    lead.post(f"/kpi/{eid}/reopen")
    with app.app_context():
        ev = db.session.get(PersonalEvaluation, eid)
        assert ev.status == EvalStatus.REVIEWED and ev.final_rating is None


def test_submit_blocked_without_department(app):
    with app.app_context():
        create_user("loner", "Loner", PW, role=Role.CC, approved=True)
    c = client(app, "loner")
    eid = start(app, c)
    c.post(f"/kpi/{eid}/save", json=plan_payload(app))
    html = c.post(f"/kpi/{eid}/plan/submit", follow_redirects=True).get_data(as_text=True)
    assert "chưa được gán Phòng, Ban" in html and plan_status(app, eid) == "draft"


def test_chain_of_approval(app):
    c, eid = to_submitted(app, who="leader_p1")             # Phó Trưởng Ban: lãnh đạo Ban nhận xét, PBT duyệt
    client(app, "leader_p1").post(f"/kpi/{eid}/approve", data=APPROVE)
    assert status(app, eid) == EvalStatus.SELF_SCORED
    assert client(app, "manager2").post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT"}).status_code == 403
    client(app, "manager1").post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT"})
    assert status(app, eid) == EvalStatus.REVIEWED
    assert client(app, "leader").post(f"/kpi/{eid}/approve", data=APPROVE).status_code == 302
    assert status(app, eid) == EvalStatus.APPROVED
    # Phó TB Thường trực: Phó Bí thư duyệt thẳng, Ban không nhận xét
    c2, e2 = to_submitted(app, who="manager1")
    assert client(app, "manager2").post(f"/kpi/{e2}/review", data={"proposed_rating": "HTT"}).status_code == 403
    client(app, "leader").post(f"/kpi/{e2}/approve", data=APPROVE)
    assert status(app, e2) == EvalStatus.APPROVED


# --------------------------- phân quyền hồ sơ --------------------------- #
def test_visibility_rules(app):
    c = client(app, "staff1")
    eid = start(app, c)
    c.post(f"/kpi/{eid}/save", json=plan_payload(app))
    for who in ("manager1", "leader", "admin", "staff2", "leader_p1"):       # đang lập: chỉ chủ hồ sơ xem
        assert client(app, who).get(f"/kpi/{eid}").status_code == 403, who
    c.post(f"/kpi/{eid}/plan/submit")
    assert client(app, "manager1").get(f"/kpi/{eid}").status_code == 200
    assert client(app, "leader").get(f"/kpi/{eid}").status_code == 200
    assert client(app, "admin").get(f"/kpi/{eid}").status_code == 200        # để mở lại Mẫu 3 khi cần
    assert client(app, "staff2").get(f"/kpi/{eid}").status_code == 403
    assert client(app, "leader_p1").get(f"/kpi/{eid}").status_code == 403    # Phó Trưởng Ban không duyệt
    assert client(app, "manager2").get(f"/kpi/{eid}").status_code == 403     # Ban khác
    assert client(app, "admin").get("/kpi/").status_code == 403


def test_staff_cannot_review_or_approve(app):
    c, eid = to_submitted(app)
    for who in ("staff2", "leader_p1"):
        other = client(app, who)
        for action in ("review", "approve", "return", "reopen", "plan/approve", "plan/return"):
            assert other.post(f"/kpi/{eid}/{action}", data={"proposed_rating": "HTT", "final_rating": "HTT", "note": "x"}).status_code == 403, (who, action)
    assert c.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT"}).status_code == 403   # tự nhận xét mình


def test_queue_scopes(app):
    to_submitted(app, who="staff1")
    to_submitted(app, who="staff2")
    q4 = ids(app)["q4"]
    c3 = client(app, "leader_p1")
    e3 = start(app, c3)
    c3.post(f"/kpi/{e3}/save", json=plan_payload(app))
    c3.post(f"/kpi/{e3}/plan/submit")                                         # Mẫu 3 chờ manager1 duyệt

    def rows(who):
        html = client(app, who).get(f"/kpi/queue?period={q4}").get_data(as_text=True)
        return {n for n in ("Staff1", "Staff2", "Leader_P1") if f"<b>{n}</b>" in html}
    assert rows("manager1") == {"Staff1", "Leader_P1"}
    assert rows("manager2") == {"Staff2"}
    assert rows("leader") == {"Staff1", "Staff2"}                             # Mẫu 3 của Leader_P1 chưa qua Ban
    assert client(app, "staff1").get("/kpi/queue").status_code == 403
    html = client(app, "manager1").get("/").get_data(as_text=True)
    assert 'class="count">2<' in html                                          # số hồ sơ chờ trên menu


def test_closed_period_locks_everything(app):
    c, eid = to_submitted(app)
    admin = client(app, "admin")
    admin.post(f"/admin/periods/{ids(app)['q4']}/toggle")
    assert c.post(f"/kpi/{eid}/withdraw").status_code == 302 and status(app, eid) == EvalStatus.SELF_SCORED
    client(app, "leader").post(f"/kpi/{eid}/approve", data=APPROVE)
    assert status(app, eid) == EvalStatus.SELF_SCORED
    assert c.post(f"/kpi/start/{ids(app)['q4']}").status_code == 302
    admin.post(f"/admin/periods/{ids(app)['q4']}/toggle")                    # mở lại
    client(app, "leader").post(f"/kpi/{eid}/approve", data=APPROVE)
    assert status(app, eid) == EvalStatus.APPROVED


def test_consistency_warnings(app):
    rows = [("Việc chưa xong", 4, 1, 100, 1, 1, 1)]
    c, eid, item_ids = ready(app, "leader_p1", rows=rows)
    d = c.post(f"/kpi/{eid}/save", json=full_payload(app, rows, item_ids=item_ids)).get_json()["data"]
    text_ = " ".join(d["warnings"])
    assert "chất lượng chấm mức 1" in text_ and "tiến độ chấm mức 1" in text_
    assert "lãnh đạo, chỉ đạo, điều hành chấm mức 1" in text_ and "dưới 100%" in text_


def test_copy_previous_quarter(app):
    c, q3, item_ids = ready(app, period="q3")
    c.post(f"/kpi/{q3}/save", json=full_payload(app, item_ids=item_ids))
    q4 = start(app, c, "q4")
    assert c.post(f"/kpi/{q4}/copy-previous").status_code == 302
    with app.app_context():
        ev = db.session.get(PersonalEvaluation, q4)
        assert len(ev.items) == 6 and all(i.actual_qty is None and i.quality_code is None for i in ev.items)
        assert not any(i.from_plan for i in ev.items) and ev.plan_status == PlanStatus.DRAFT
        assert ev.axis_results[0].max_score == 70
    c.post(f"/kpi/{q4}/copy-previous")                                          # đã có dòng -> không chép đè
    with app.app_context():
        assert len(db.session.get(PersonalEvaluation, q4).items) == 6


def test_pages_render(app):
    c, eid, item_ids = ready(app)
    c.post(f"/kpi/{eid}/save", json=full_payload(app, item_ids=item_ids))
    assert c.get("/kpi/").status_code == 200
    html = c.get(f"/kpi/{eid}").get_data(as_text=True)
    assert 'id="kpi-data"' in html and "kpi_personal.js" in html
    for fragment in ("Bước 1", "Mẫu 3", "Bước 2", "Mẫu 4", "Bước 3", "Mẫu 5", "Họ và tên", "Phòng 1"):
        assert fragment in html, fragment
    assert client(app, "admin").get("/admin/periods").status_code == 200
    assert client(app, "manager1").get(f"/kpi/queue").status_code == 200


def test_period_admin(app):
    admin = client(app, "admin")
    admin.post("/admin/periods/create", data={"year": 2027, "quarter": 1})
    admin.post("/admin/periods/create", data={"year": 2027, "quarter": 1})       # trùng
    admin.post("/admin/periods/create", data={"year": 1999, "quarter": 5})       # sai
    with app.app_context():
        assert Period.query.filter_by(year=2027).count() == 1
        p = Period.query.filter_by(year=2027).first()
        assert str(p.deadline_m1_m3) == "2027-01-15" and str(p.deadline_m2) == "2027-03-18"
        pid = p.id
    admin.post(f"/admin/periods/{pid}/deadlines", data={"deadline_m1_m3": "2027-01-10", "deadline_m2": "2027-03-15", "deadline_m4_m8": "x"})
    with app.app_context():
        assert str(db.session.get(Period, pid).deadline_m1_m3) == "2027-01-15"   # ngày sai -> không đổi gì
    assert client(app, "manager1").post("/admin/periods/create", data={"year": 2028, "quarter": 1}).status_code == 403


# --------------------------- KPI tập thể --------------------------- #
def col_payload(app, with_result=True):
    axis = ids(app)["axis"][1]
    mgr = uid(app, "manager1")
    return {
        "other_key_tasks": "Khác", "innovative_models": "Mô hình",
        "indicators": [
            {"axis_id": axis, "task_group": "thuong_xuyen", "content": "Kết nạp đảng viên", "unit": "người",
             "annual_target": 100, "quarter_target": 20, "deadline_date": "2026-12-30", "leader_id": mgr,
             "result": 25 if with_result else None, "actual_date": "2026-12-29" if with_result else ""},
            {"axis_id": axis, "task_group": "trong_tam", "content": "Đề án", "quarter_target": 1,
             "result": 0 if with_result else None},
        ],
    }


def col_url(app, dept="p1", period="q4", tail=""):
    i = ids(app)
    return f"/collective/{i[period]}/{i[dept]}{tail}"


def col_status(app, dept="p1"):
    with app.app_context():
        return CollectiveEvaluation.query.filter_by(department_id=ids(app)[dept]).first().status


def col_results(app, rows, with_result=True):
    """Payload Mẫu 2: chỉ cột kết quả cho các dòng của Mẫu 1."""
    vals = [(25, "2026-12-29"), (0, "")]
    return {"other_key_tasks": "Khác", "innovative_models": "Mô hình",
            "indicators": [{"id": r["id"], "result": vals[n][0] if with_result else None,
                            "actual_date": vals[n][1] if with_result else ""} for n, r in enumerate(rows)]}


def m1_approved(app):
    """Phó TB TT lập, trình Mẫu 1; Phó Bí thư duyệt. Trả về (client đơn vị, các dòng)."""
    mgr = client(app, "manager1")
    mgr.get(col_url(app))
    rows = mgr.post(col_url(app, tail="/save"), json=col_payload(app)).get_json()["data"]["indicators"]
    mgr.post(col_url(app, tail="/submit"))
    client(app, "leader").post(col_url(app, tail="/approve"))
    assert col_status(app) == "approved"
    return mgr, rows


def col_result_status(app, dept="p1"):
    with app.app_context():
        return CollectiveEvaluation.query.filter_by(department_id=ids(app)[dept]).first().result_status


def test_collective_flow(app):
    mgr = client(app, "manager1")
    assert mgr.get("/collective/").status_code == 302                          # chỉ có một Ban -> vào thẳng
    html = mgr.get(col_url(app)).get_data(as_text=True)
    assert "Phòng 1" in html and "Mẫu 2 mở sau khi Mẫu 1 được duyệt" in html
    d = mgr.post(col_url(app, tail="/save"), json=col_payload(app)).get_json()["data"]
    assert d["indicators"][0]["result"] is None                                 # Mẫu 1 không nhận kết quả
    p = col_payload(app)
    p["indicators"][1]["quarter_target"] = None
    mgr.post(col_url(app, tail="/save"), json=p)
    mgr.post(col_url(app, tail="/submit"))
    assert col_status(app) == "draft"                                           # thiếu chỉ tiêu quý
    mgr.post(col_url(app, tail="/save"), json=col_payload(app))
    mgr.post(col_url(app, tail="/submit"))
    assert col_status(app) == "submitted"
    assert mgr.post(col_url(app, tail="/save"), json=col_payload(app)).status_code == 409
    assert mgr.post(col_url(app, tail="/approve")).status_code == 403            # đơn vị không tự duyệt
    mgr.post(col_url(app, tail="/withdraw"))
    assert col_status(app) == "draft"
    mgr.post(col_url(app, tail="/submit"))
    lead = client(app, "leader")
    lead.post(col_url(app, tail="/return"), data={"note": "Kiểm tra lại chỉ tiêu"})
    assert col_status(app) == "draft"
    assert "Kiểm tra lại chỉ tiêu" in mgr.get(col_url(app)).get_data(as_text=True)
    mgr.post(col_url(app, tail="/submit"))
    lead.post(col_url(app, tail="/approve"))
    assert col_status(app) == "approved" and col_result_status(app) == "draft"
    with app.app_context():
        assert CollectiveEvaluation.query.first().confirmed_by.username == "leader"


def test_mau2_locked_until_mau1_approved_and_rows_carried(app):
    mgr = client(app, "manager1")
    mgr.get(col_url(app))
    rows = mgr.post(col_url(app, tail="/save"), json=col_payload(app)).get_json()["data"]["indicators"]
    assert mgr.post(col_url(app, tail="/results/save"), json=col_results(app, rows)).status_code == 409
    mgr.post(col_url(app, tail="/submit"))
    assert mgr.post(col_url(app, tail="/results/save"), json=col_results(app, rows)).status_code == 409
    client(app, "leader").post(col_url(app, tail="/approve"))
    html = mgr.get(col_url(app)).get_data(as_text=True)
    assert "Mẫu 2 mở sau khi Mẫu 1 được duyệt" not in html and "Trình Mẫu 2" in html
    r = mgr.post(col_url(app, tail="/results/save"), json=col_results(app, rows))
    d = r.get_json()["data"]
    a, b = d["indicators"]                                                      # dòng lấy từ Mẫu 1
    assert (a["content"], a["quarter_target"]) == ("Kết nạp đảng viên", 20)
    assert (a["pct_quarter"], a["pct_year"], a["on_time"]) == (125.0, 25.0, True)
    assert b["pct_quarter"] == 0.0
    assert d["summary"] == {"total": 2, "with_result": 2, "achieved": 1, "on_time": 1, "avg_pct": 50.0}
    bad = col_results(app, rows)
    bad["indicators"][0]["id"] = 99999                                          # dòng không có trong Mẫu 1
    assert mgr.post(col_url(app, tail="/results/save"), json=bad).status_code == 400
    assert mgr.post(col_url(app, tail="/save"), json=col_payload(app)).status_code == 409   # Mẫu 1 đã duyệt: khóa


def test_mau2_separate_submit_approve_and_reopen_mau1(app):
    mgr, rows = m1_approved(app)
    lead = client(app, "leader")
    mgr.post(col_url(app, tail="/results/save"), json=col_results(app, rows, with_result=False))
    mgr.post(col_url(app, tail="/results/submit"))
    assert col_result_status(app) == "draft"                                    # thiếu kết quả
    mgr.post(col_url(app, tail="/results/save"), json=col_results(app, rows))
    mgr.post(col_url(app, tail="/results/submit"))
    assert col_result_status(app) == "submitted"
    assert mgr.post(col_url(app, tail="/results/approve")).status_code == 403
    lead.post(col_url(app, tail="/results/return"), data={"note": "Bổ sung minh chứng"})
    assert col_result_status(app) == "draft"
    mgr.post(col_url(app, tail="/results/submit"))
    lead.post(col_url(app, tail="/results/approve"), data={"rating": "HTT"})
    assert col_result_status(app) == "approved" and col_status(app) == "approved"
    with app.app_context():
        ce = CollectiveEvaluation.query.first()
        assert ce.rating == "HTT" and ce.result_confirmed_by.username == "leader"
    lead.post(col_url(app, tail="/results/reopen"))
    assert col_result_status(app) == "draft"
    # mở lại Mẫu 1: Mẫu 2 khóa, kết quả vẫn gắn với các dòng cũ
    lead.post(col_url(app, tail="/reopen"))
    assert col_status(app) == "draft"
    html = mgr.get(col_url(app)).get_data(as_text=True)
    assert "Mẫu 2 mở sau khi Mẫu 1 được duyệt" in html
    assert mgr.post(col_url(app, tail="/results/save"), json=col_results(app, rows)).status_code == 409
    keep = col_payload(app)
    for n, r in enumerate(rows):
        keep["indicators"][n]["id"] = r["id"]
    d = mgr.post(col_url(app, tail="/save"), json=keep).get_json()["data"]
    assert d["indicators"][0]["result"] == 25                                   # kết quả giữ nguyên
    mgr.post(col_url(app, tail="/submit"))
    lead.post(col_url(app, tail="/approve"))
    assert col_status(app) == "approved" and col_result_status(app) == "draft"
    dash = lead.get("/").get_data(as_text=True)
    assert "Mẫu 1 · Danh mục chỉ tiêu" in dash and "Mẫu 2 · Kết quả thực hiện" in dash


def test_upgrade_marks_old_approved_collective(app):
    mgr, rows = m1_approved(app)
    with app.app_context():
        from models import CollectiveKPI
        for r in rows:
            db.session.add(CollectiveKPI(indicator_id=r["id"], result=1))
        db.session.execute(text("ALTER TABLE collective_evaluations DROP COLUMN result_status"))
        db.session.commit()
        upgrade_schema()
    assert col_status(app) == "approved" and col_result_status(app) == "approved"


def test_collective_only_ptbtt_edits(app):
    mgr = client(app, "manager1")
    mgr.get(col_url(app))                                                       # tạo hồ sơ Phòng 1
    for who in ("staff1", "leader_p1", "leader", "admin"):                      # chỉ xem
        c = client(app, who)
        assert c.get(col_url(app)).status_code == 200, who
        assert c.post(col_url(app, tail="/save"), json=col_payload(app)).status_code == 403, who
    s2 = client(app, "staff2")
    assert s2.get(col_url(app)).status_code == 403                               # Ban khác
    assert "chưa lập KPI tập thể" in s2.get(col_url(app, "p2")).get_data(as_text=True)
    assert "chưa lập KPI tập thể" in client(app, "leader").get(col_url(app, "p2")).get_data(as_text=True)
    with app.app_context():
        assert CollectiveEvaluation.query.filter_by(department_id=ids(app)["p2"]).count() == 0
    assert client(app, "manager1").get(col_url(app, "p2")).status_code == 403
    assert client(app, "manager2").get(col_url(app, "p2")).status_code == 200    # Phó TB TT Phòng 2 lập
    with app.app_context():
        assert CollectiveEvaluation.query.filter_by(department_id=ids(app)["p2"]).count() == 1
    assert client(app, "manager2").post(col_url(app, tail="/save"), json=col_payload(app)).status_code == 403


def test_collective_validation_and_ownership(app):
    mgr = client(app, "manager1")
    mgr.get(col_url(app))
    p = col_payload(app)
    p["indicators"][0]["leader_id"] = uid(app, "staff2")                         # không thuộc Ban này
    assert mgr.post(col_url(app, tail="/save"), json=p).status_code == 400
    p = col_payload(app)
    p["indicators"][0]["quarter_target"] = -5
    assert mgr.post(col_url(app, tail="/save"), json=p).status_code == 400
    p = col_payload(app)
    mgr.post(col_url(app, tail="/save"), json=p)
    with app.app_context():
        other = Indicator.query.first().id
    mgr2 = client(app, "manager2")
    mgr2.get(col_url(app, "p2"))
    p2 = col_payload(app)
    p2["indicators"][0]["id"] = other                                           # dòng của Phòng 1
    p2["indicators"][0]["leader_id"] = None
    assert mgr2.post(col_url(app, "p2", tail="/save"), json=p2).status_code == 400


def test_deleting_linked_indicator_unlinks_personal_item(app):
    mgr = client(app, "manager1")
    mgr.get(col_url(app))
    d = mgr.post(col_url(app, tail="/save"), json=col_payload(app)).get_json()["data"]
    ind_id = d["indicators"][0]["id"]
    s = client(app, "staff1")
    eid = start(app, s)
    p = plan_payload(app, rows=[EXAMPLE[0]])
    p["items"][0]["indicator_id"] = ind_id
    assert s.post(f"/kpi/{eid}/save", json=p).status_code == 200
    html = s.get(f"/kpi/{eid}").get_data(as_text=True)
    import json, re
    page = json.loads(re.search(r'<script id="kpi-data" type="application/json">(.*?)</script>', html, re.S).group(1))
    assert [x["content"] for x in page["indicators"]] == ["Kết nạp đảng viên", "Đề án"]   # danh sách chọn nhanh từ chỉ tiêu đơn vị
    keep = col_payload(app)
    keep["indicators"] = [dict(d["indicators"][1])]
    assert mgr.post(col_url(app, tail="/save"), json=keep).status_code == 200
    with app.app_context():
        assert db.session.get(PersonalEvaluation, eid).items[0].indicator_id is None
    p["items"][0]["indicator_id"] = ind_id                                       # chỉ tiêu đã xóa -> báo lỗi
    assert s.post(f"/kpi/{eid}/save", json=p).status_code == 400


# --------------------------- nâng cấp CSDL --------------------------- #
def test_upgrade_schema_adds_missing_columns_keeps_data(app):
    with app.app_context():
        db.session.execute(text("ALTER TABLE personal_evaluations DROP COLUMN reviewer_score"))
        db.session.execute(text("ALTER TABLE personal_evaluations DROP COLUMN use_leadership"))
        db.session.execute(text("ALTER TABLE users DROP COLUMN session_token"))
        db.session.execute(text("DROP TABLE workflow_logs"))
        db.session.commit()
        changes = upgrade_schema()
        assert {"personal_evaluations.reviewer_score", "personal_evaluations.use_leadership",
                "users.session_token", "workflow_logs"} <= set(changes) | {"workflow_logs"}
        tokens = [u.session_token for u in User.query.all()]
        assert all(len(t) == 64 for t in tokens) and len(set(tokens)) == len(tokens)   # điền token riêng cho từng người
        assert upgrade_schema() == []                                                   # chạy lại không đổi gì
        assert User.query.count() == 7


def test_old_session_survives_upgrade_login_still_works(app):
    with app.app_context():
        db.session.execute(text("ALTER TABLE personal_evaluations DROP COLUMN use_leadership"))
        db.session.commit()
        upgrade_schema()
    c = client(app, "staff1")
    assert start(app, c) > 0


def test_others_cannot_see_draft_scores(app):
    c, eid, item_ids = ready(app)
    c.post(f"/kpi/{eid}/save", json=full_payload(app, item_ids=item_ids))
    import json, re
    html = client(app, "leader").get(f"/kpi/{eid}").get_data(as_text=True)     # Mẫu 3 đã duyệt, Mẫu 5 chưa trình
    data = json.loads(re.search(r'<script id="kpi-data" type="application/json">(.*?)</script>', html, re.S).group(1))
    assert data["eval"]["items"] == [] and len(data["eval"]["plan"]) == 6 and "chưa trình Mẫu 5" in html
    assert len(page_data(c, eid)["eval"]["items"]) == 6                         # chủ hồ sơ vẫn thấy đủ


def test_upgrade_marks_old_submitted_profiles_plan_approved(app):
    c, eid = to_submitted(app)
    with app.app_context():
        db.session.execute(text("ALTER TABLE personal_evaluations DROP COLUMN plan_status"))
        db.session.execute(text("ALTER TABLE personal_kpi_items DROP COLUMN from_plan"))
        db.session.commit()
        upgrade_schema()
    assert plan_status(app, eid) == PlanStatus.APPROVED
    with app.app_context():
        assert all(i.from_plan for i in db.session.get(PersonalEvaluation, eid).items)


# --------------------------- vai trò bổ sung (điều chỉnh lần 3) --------------------------- #
NEW_LEADERS = (Role.PCVPTT, Role.PCNTT)
NEW_DEPUTIES = (Role.PCVP, Role.PCN, Role.UVCT)


def test_new_unit_leader_roles_review_and_edit_collective(app):
    for n, role in enumerate(NEW_LEADERS):
        with app.app_context():
            u = User.query.filter_by(username="manager1").first()
            u.role = role
            db.session.commit()
        # Mẫu 3 của công chức cùng đơn vị: lãnh đạo đơn vị duyệt cấp đơn vị
        c = client(app, f"staff1")
        eid = start(app, c, "q4" if n == 0 else "q3")
        c.post(f"/kpi/{eid}/save", json=plan_payload(app))
        c.post(f"/kpi/{eid}/plan/submit")
        assert client(app, "manager1").post(f"/kpi/{eid}/plan/approve").status_code == 302, role
        assert plan_status(app, eid) == PlanStatus.BAN_APPROVED, role
        # KPI tập thể của đơn vị mình
        mgr = client(app, "manager1")
        mgr.get(col_url(app, period="q4" if n == 0 else "q3"))
        assert mgr.post(col_url(app, period="q4" if n == 0 else "q3", tail="/save"),
                        json=col_payload(app)).status_code == 200, role


def test_new_deputy_roles_have_d_and_cannot_review(app):
    to_submitted(app, "staff1")
    with app.app_context():
        eid = PersonalEvaluation.query.first().id
    for role in NEW_DEPUTIES:
        with app.app_context():
            u = User.query.filter_by(username="leader_p1").first()
            u.role = role
            db.session.commit()
            assert u.has_leadership_kpi
        c = client(app, "leader_p1")
        assert c.post(f"/kpi/{eid}/review", data={"proposed_rating": "HTT"}).status_code == 403, role
        assert c.get("/kpi/queue").status_code == 403, role
        assert c.post(col_url(app, tail="/save"), json=col_payload(app)).status_code == 403, role
    _, e2, _ = ready(app, "leader_p1")                 # Ủy viên CT UBKT: Mẫu 3 do lãnh đạo đơn vị duyệt, có D
    assert page_data(client(app, "leader_p1"), e2)["leadership"] is True


def test_role_order_in_selects():
    order = [c for c, _ in Role.CHOICES]
    assert order == [Role.ADMIN, Role.PBT, Role.PTBTT, Role.PCVPTT, Role.PCNTT,
                     Role.PTB, Role.PCVP, Role.PCN, Role.UVCT, Role.CC]


# --------------------------- bố cục biểu mẫu, báo cáo Mẫu 6, 7, 8 --------------------------- #
def test_mau5_axis_goals_results_saved(app):
    c, eid, item_ids = ready(app)
    p = full_payload(app, item_ids=item_ids)
    p["axes"][0].update(goals="Mục tiêu trục 4", results="Đã hoàn thành", note="Ghi chú")
    d = c.post(f"/kpi/{eid}/save", json=p).get_json()["data"]
    assert (d["axes"][0]["goals"], d["axes"][0]["results"], d["axes"][0]["note"]) == ("Mục tiêu trục 4", "Đã hoàn thành", "Ghi chú")
    html = c.get(f"/kpi/{eid}").get_data(as_text=True)
    for fragment in ("NHIỆM VỤ ĐỀ RA KẾ HOẠCH VÀ THỰC TẾ TRIỂN KHAI", "BẢN TỰ ĐÁNH GIÁ, XẾP LOẠI CỦA CÁ NHÂN",
                     "A. NHÓM TIÊU CHÍ CHUNG (30 ĐIỂM)", "B. KẾT QUẢ THỰC HIỆN NHIỆM VỤ ĐƯỢC GIAO (70 ĐIỂM)",
                     "II. Tự đề xuất xếp loại", "III. Nhận xét, đánh giá của cấp có thẩm quyền", "Chức vụ Đảng"):
        assert fragment in html, fragment


def test_reports_m6_m7_m8(app):
    c, eid = to_submitted(app, "staff1")
    with app.app_context():
        from models import PositionGroup
        g = PositionGroup.query.filter_by(code="M8-PHONG-PHO").first()
        u = User.query.filter_by(username="staff1").first()
        u.position_group_id = g.id
        ev = db.session.get(PersonalEvaluation, eid)
        ev.position_group_id = g.id
        db.session.commit()
    lead = client(app, "leader")
    q4 = ids(app)["q4"]
    m6 = lead.get(f"/reports/m6?period={q4}").get_data(as_text=True)
    assert "TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ VÀ ĐỀ XUẤT XẾP LOẠI" in m6 and "Staff1" in m6 and "Hoàn thành tốt nhiệm vụ" in m6
    lead.post(f"/kpi/{eid}/approve", data={"final_rating": "HTXS", "rating_reason": "Vượt mức", "response_level": "1"})
    m8 = lead.get(f"/reports/m8?period={q4}").get_data(as_text=True)
    assert "CẤP THÀNH PHỐ" in m8 and "Cấp phó trưởng phòng, ban" in m8 and "HTTNV trở lên" in m8
    assert "Tính 1 hồ sơ" in m8 and "100.00" in m8                 # 1 HTXS / 1 HTTNV trở lên
    m7 = lead.get(f"/reports/m7?period={q4}").get_data(as_text=True)
    assert "CẤP XÃ VÀ TƯƠNG ĐƯƠNG" in m7 and "Tính 0 hồ sơ" in m7
    assert client(app, "manager2").get(f"/reports/m6?period={q4}").get_data(as_text=True).count("Staff1") == 0
    assert client(app, "staff1").get("/reports/m6").status_code == 403
    assert client(app, "admin").get("/reports/m9").status_code == 404


# --------------------------- xuất Excel --------------------------- #
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(resp):
    from io import BytesIO
    from openpyxl import load_workbook
    assert resp.status_code == 200, resp.status_code
    assert resp.mimetype == XLSX
    return load_workbook(BytesIO(resp.data)).active


def texts(ws):
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]


def test_export_personal_m3_m4_m5(app):
    c, eid = to_submitted(app, "leader_p1")
    ws = workbook(c.get(f"/kpi/{eid}/export/m3"))
    t = texts(ws)
    assert "DANH MỤC SẢN PHẨM/CÔNG VIỆC CỦA CÁ NHÂN CẦN THỰC HIỆN" in ws["A3"].value and "Mẫu 3" in t
    assert [f"({i})" for i in range(1, 10)] == [x for x in t if x.startswith("(") and x[1:-1].isdigit()][:9]
    assert "Báo cáo tháng" in t
    ws = workbook(c.get(f"/kpi/{eid}/export/m4"))
    t = texts(ws)
    assert "NHIỆM VỤ ĐỀ RA KẾ HOẠCH VÀ THỰC TẾ TRIỂN KHAI" in ws["A3"].value
    assert "THỰC TẾ TRIỂN KHAI" in t and "Hệ số KPI vượt" in t and "(19)" in t and "TỔNG" in t
    assert ws.page_setup.orientation == "landscape" and ws.print_title_rows
    ws = workbook(c.get(f"/kpi/{eid}/export/m5"))
    t = texts(ws)
    assert "BẢN TỰ ĐÁNH GIÁ, XẾP LOẠI CỦA CÁ NHÂN" in ws["A3"].value
    assert any("NHÓM TIÊU CHÍ CHUNG (30 ĐIỂM)" in x for x in t) and "Tổng (A) =" in t and "TỔNG (A+ B) = " in t
    assert any(x == "97.37" or x.startswith("97.37") for x in t)                       # tổng điểm ví dụ Phụ lục 2
    resp = c.get(f"/kpi/{eid}/export/m4")
    assert "Mau4_leader_p1_Q4-2026.xlsx" in resp.headers["Content-Disposition"]
    # quyền: người không xem được hồ sơ thì không xuất được
    assert client(app, "staff2").get(f"/kpi/{eid}/export/m3").status_code == 403
    assert client(app, "manager2").get(f"/kpi/{eid}/export/m4").status_code == 403
    assert workbook(client(app, "manager1").get(f"/kpi/{eid}/export/m5"))                # lãnh đạo đơn vị xem được
    assert c.get(f"/kpi/{eid}/export/m9").status_code == 404


def test_export_personal_hides_drafts(app):
    c, eid, item_ids = ready(app)
    c.post(f"/kpi/{eid}/save", json=full_payload(app, item_ids=item_ids))
    assert workbook(c.get(f"/kpi/{eid}/export/m4"))                                     # chủ hồ sơ xuất được bản nháp
    lead = client(app, "leader")
    assert workbook(lead.get(f"/kpi/{eid}/export/m3"))                                  # Mẫu 3 đã duyệt
    assert lead.get(f"/kpi/{eid}/export/m4").status_code == 403                         # Mẫu 4, 5 chưa trình
    assert lead.get(f"/kpi/{eid}/export/m5").status_code == 403
    c2 = client(app, "staff2")
    e2 = start(app, c2)
    assert c2.get(f"/kpi/{e2}/export/m4").status_code == 404                            # chưa tới bước Mẫu 4


def test_export_collective_m1_m2(app):
    mgr, rows = m1_approved(app)
    mgr.post(col_url(app, tail="/results/save"), json=col_results(app, rows))
    ws = workbook(mgr.get(col_url(app, tail="/export/m1")))
    assert "DANH MỤC CHỈ TIÊU" in ws["A3"].value and "(10)" in texts(ws) and "Kết nạp đảng viên" in texts(ws)
    ws = workbook(client(app, "leader").get(col_url(app, tail="/export/m2")))
    t = texts(ws)
    assert "KẾT QUẢ THỰC HIỆN NHIỆM VỤ CỦA PHÒNG 1" in ws["A3"].value
    assert "(8)=(7)/(6)*100" in t and "125" in " ".join(t) and "Khác" in t
    assert client(app, "staff2").get(col_url(app, tail="/export/m1")).status_code == 403
    m2 = client(app, "manager2")
    m2.get(col_url(app, "p2"))
    assert m2.get(col_url(app, "p2", tail="/export/m2")).status_code == 404             # Mẫu 1 chưa duyệt
    assert client(app, "leader").get(col_url(app, "p1", period="q3", tail="/export/m1")).status_code == 404


def test_export_reports(app):
    c, eid = to_submitted(app, "staff1")
    lead = client(app, "leader")
    q4 = ids(app)["q4"]
    ws = workbook(lead.get(f"/reports/m6/export?period={q4}"))
    assert "TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ" in ws["A3"].value and "Staff1" in texts(ws) and "(7)" in texts(ws)
    for form, title in (("m7", "CẤP XÃ"), ("m8", "CẤP THÀNH PHỐ")):
        ws = workbook(lead.get(f"/reports/{form}/export?period={q4}"))
        assert title in ws["A3"].value and "HTTNV trở lên" in texts(ws)
    assert client(app, "staff1").get(f"/reports/m6/export?period={q4}").status_code == 403
    assert lead.get("/reports/m9/export").status_code == 404
